"""End-to-end test of the ingest api: templates, uploads (good, bad, refused) and the reject list.

Runs against a live server over plain http (standard library only, no pytest / requests needed):

    python -m uvicorn app.main:app --port 8765          # in one terminal, from Backend/
    python tests/test_ingest_api.py --reset             # in another, from Backend/

--reset empties the ingest tables (raw_*, ingest_rejects, ingest_files) and Backend/uploads first, because a
file that is already loaded is refused as a duplicate. Test data: a folder with Correct/ and Dummy/ (broken copies).
"""

import argparse
import io
import json
import sys
import time
import uuid
import urllib.error
import urllib.request
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))             # Backend/ on the path
from openpyxl import load_workbook                                       # noqa: E402

from app.core.database import get_postgres_cursor                         # noqa: E402
from app.ingest.registry import FILE_SPECS                                # noqa: E402

UPLOAD_DIR = Path(__file__).resolve().parents[1] / "uploads"
FILES = {"bom_extract": "BOM Extract.xlsx",
         "shelf_life": "FG_Shelf_Life_days_QMS.xlsx",
         "cycle_time": "PPS Products Cycle Time 05-02-2026.xlsx"}
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

results: list[tuple[bool, str, str]] = []


def check(name: str, ok: bool, detail: str = "") -> bool:
    results.append((ok, name, detail))
    print(f"  {'PASS' if ok else 'FAIL'}  {name}" + (f"  - {detail}" if detail and not ok else ""))
    return ok


# ---------------------------------------------------------------- http, standard library only

def request(method: str, url: str, body: bytes | None = None, headers: dict | None = None):
    req = urllib.request.Request(url, data=body, method=method, headers=headers or {})
    try:
        with urllib.request.urlopen(req, timeout=600) as r:
            return r.status, dict(r.headers), r.read()
    except urllib.error.HTTPError as e:
        return e.code, dict(e.headers), e.read()


def as_json(raw: bytes) -> Any:
    try:
        return json.loads(raw)
    except ValueError:
        return raw.decode(errors="replace")


def upload(base: str, content: bytes | None, filename: str, fields: dict):
    boundary = uuid.uuid4().hex
    parts = []
    for k, v in fields.items():
        parts.append(f'--{boundary}\r\nContent-Disposition: form-data; name="{k}"\r\n\r\n{v}\r\n'.encode())
    if content is not None:
        parts.append(f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="{filename}"\r\n'
                     f"Content-Type: {XLSX}\r\n\r\n".encode() + content + b"\r\n")
    body = b"".join(parts) + f"--{boundary}--\r\n".encode()
    t = time.time()
    status, _, raw = request("POST", f"{base}/ingest/upload", body,
                             {"Content-Type": f"multipart/form-data; boundary={boundary}"})
    return status, as_json(raw), time.time() - t


# ---------------------------------------------------------------- database, read-only except --reset

def db(sql: str, params=()):
    conn, cur = get_postgres_cursor()
    try:
        cur.execute(sql, params)
        rows = cur.fetchall()
        conn.rollback()
        return rows
    finally:
        conn.close()


def reset():
    conn, cur = get_postgres_cursor()
    try:
        cur.execute("TRUNCATE raw_bom_extract, raw_shelf_life, raw_cycle_time, ingest_rejects, ingest_files "
                    "RESTART IDENTITY")
        conn.commit()
    finally:
        conn.close()
    for f in UPLOAD_DIR.glob("*"):
        if f.is_file():
            f.unlink()


def current_file(file_type: str):
    rows = db("SELECT file_id FROM ingest_files WHERE file_type = %s AND is_current", (file_type,))
    return rows[0][0] if len(rows) == 1 else None


# ---------------------------------------------------------------- the test

def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--base-url", default="http://127.0.0.1:8765")
    ap.add_argument("--data", default=r"C:\Users\yashb\Downloads\inget data")
    ap.add_argument("--reset", action="store_true", help="empty the ingest tables before testing")
    a = ap.parse_args()
    base, good, bad = a.base_url.rstrip("/"), Path(a.data) / "Correct", Path(a.data) / "Dummy"

    if a.reset:
        reset()
        print("ingest tables and uploads/ emptied\n")

    # ------------------------------------------------ templates
    print("GET /ingest/template/{file_type}")
    for key, spec in FILE_SPECS.items():
        status, headers, raw = request("GET", f"{base}/ingest/template/{key}")
        if check(f"{key}: 200 + xlsx", status == 200 and headers.get("content-type", "").startswith(XLSX),
                 f"status {status}, {headers.get('content-type')}"):
            wb = load_workbook(io.BytesIO(raw))
            check(f"{key}: sheets Template + Guidelines", wb.sheetnames == ["Template", "Guidelines"], str(wb.sheetnames))
            heads = [c.value for c in wb.worksheets[0][1]]
            check(f"{key}: headers match the spec", heads == [c.source for c in spec.columns])
            check(f"{key}: download file name", f"{key}_template.xlsx" in headers.get("content-disposition", ""),
                  headers.get("content-disposition", ""))
    status, _, raw = request("GET", f"{base}/ingest/template/po_receipts")
    check("unknown type refused (4xx)", 400 <= status < 500, f"status {status}")

    # ------------------------------------------------ correct files
    print("\nPOST /ingest/upload - correct files")
    expected_rows = {"bom_extract": 50927, "shelf_life": 803, "cycle_time": 148}
    good_ids = {}
    for key, name in FILES.items():
        status, body, secs = upload(base, (good / name).read_bytes(), name,
                                    {"uploaded_by": "api-test", "file_type": key})
        ok = status == 201 and isinstance(body, dict) and body.get("status") == "published"
        check(f"{key}: 201 published ({secs:.1f}s)", ok, f"status {status}: {str(body)[:200]}")
        if ok:
            good_ids[key] = body["file_id"]
            check(f"{key}: {expected_rows[key]} rows loaded, 0 rejected",
                  body.get("rows_loaded") == expected_rows[key] and body.get("rows_rejected") == 0,
                  f"loaded {body.get('rows_loaded')}, rejected {body.get('rows_rejected')}")
            check(f"{key}: is current", body.get("is_current") is True)
    check("cycle_time: plant recorded as the scope",
          db("SELECT period_key FROM ingest_files WHERE file_id = %s", (good_ids.get("cycle_time", -1),))
          == [("PSM - Thervoykandigai MFG",)])
    check("cycle_time: code check empty", db("SELECT count(*) FROM v_cycle_time_code_check") == [(0,)])

    # ------------------------------------------------ refused before loading
    print("\nPOST /ingest/upload - refused")
    name = FILES["shelf_life"]
    status, body, _ = upload(base, (good / name).read_bytes(), name, {"uploaded_by": "api-test", "file_type": "shelf_life"})
    check("same file again -> 409 with the earlier file_id",
          status == 409 and isinstance(body, dict) and body.get("detail", {}).get("file_id") == good_ids.get("shelf_life"),
          f"status {status}: {str(body)[:200]}")

    status, body, _ = upload(base, b"plain text, not a workbook", "fake.xlsx", {"uploaded_by": "api-test", "file_type": "shelf_life"})
    check("not an xlsx (renamed text) -> 400", status == 400, f"status {status}: {str(body)[:200]}")

    status, body, _ = upload(base, b"a,b\n1,2\n", "data.csv", {"uploaded_by": "api-test", "file_type": "shelf_life"})
    check("a .csv -> 400", status == 400, f"status {status}: {str(body)[:200]}")
    check("  message says how to fix it (Save As .xlsx)", "Save As" in str(body), str(body)[:160])

    status, body, _ = upload(base, (good / name).read_bytes(), name, {"uploaded_by": "api-test", "file_type": "po_receipts"})
    check("unknown file_type -> 422", status == 422, f"status {status}")

    status, body, _ = upload(base, (good / name).read_bytes(), name, {"file_type": "shelf_life"})
    check("uploaded_by missing -> 422", status == 422, f"status {status}")

    status, body, _ = upload(base, (good / name).read_bytes(), name, {"uploaded_by": "", "file_type": "shelf_life"})
    check("uploaded_by empty -> refused (4xx)", 400 <= status < 500, f"status {status}: {str(body)[:160]}")

    status, body, _ = upload(base, (good / name).read_bytes(), name, {"uploaded_by": "   ", "file_type": "shelf_life"})
    check("uploaded_by only spaces -> 422", status == 422, f"status {status}: {str(body)[:160]}")

    status, body, _ = upload(base, None, "", {"uploaded_by": "api-test", "file_type": "shelf_life"})
    check("no file -> 422", status == 422, f"status {status}")

    # ------------------------------------------------ the file is read and refused
    print("\nPOST /ingest/upload - failed (file-level problem)")
    name = FILES["shelf_life"]
    status, body, _ = upload(base, (good / name).read_bytes(), name, {"uploaded_by": "api-test", "file_type": "bom_extract"})
    ok = status == 201 and isinstance(body, dict) and body.get("status") == "failed"
    check("shelf life file sent as bom_extract -> failed", ok, f"status {status}: {str(body)[:200]}")
    if ok:
        check("  names the missing columns", "ORGANIZATION_CODE" in (body.get("error") or ""), body.get("error", "")[:120])

    name = FILES["bom_extract"]
    status, body, _ = upload(base, (bad / name).read_bytes(), name, {"uploaded_by": "api-test", "file_type": "bom_extract"})
    ok = status == 201 and isinstance(body, dict) and body.get("status") == "failed"
    check("BOM without the BASIS_TYPE column -> failed", ok, f"status {status}: {str(body)[:200]}")
    if ok:
        check("  says BASIS_TYPE is missing", "BASIS_TYPE" in (body.get("error") or ""), body.get("error", "")[:120])
    check("  the good BOM is still current", current_file("bom_extract") == good_ids.get("bom_extract"))

    # a workbook last saved by a script: formulas without results
    wb = load_workbook(good / FILES["cycle_time"])
    buf = io.BytesIO(); wb.save(buf)
    status, body, _ = upload(base, buf.getvalue(), "cycle_saved_by_script.xlsx", {"uploaded_by": "api-test", "file_type": "cycle_time"})
    ok = status == 201 and isinstance(body, dict) and body.get("status") == "failed"
    check("cycle time saved by a script -> failed", ok, f"status {status}: {str(body)[:200]}")
    if ok:
        check("  tells the user to save it in Excel", "Excel" in (body.get("error") or ""), body.get("error", "")[:120])

    # ------------------------------------------------ rejected: row problems
    print("\nPOST /ingest/upload - rejected (row problems)")
    rejected = {}
    for key, rows, columns in [("shelf_life", [5, 10, 15, 20], {"SHELF_LIFE_DAYS", "Itemcode"}),
                               ("cycle_time", [40, 60, 90, 120], {"MAX BATCH SIZE (IN KGS)", "PACKING SIZE",
                                                                  "CYCLE TIME EQUIPMENT WITH CLEANING(IN HRS)",
                                                                  "Equipment ID"})]:
        name = FILES[key]
        status, body, _ = upload(base, (bad / name).read_bytes(), name, {"uploaded_by": "api-test", "file_type": key})
        ok = status == 201 and isinstance(body, dict) and body.get("status") == "rejected"
        check(f"{key}: rejected", ok, f"status {status}: {str(body)[:200]}")
        if not ok:
            continue
        rejected[key] = body["file_id"]
        problems = body.get("problems") or []
        check(f"{key}: response lists the problems", len(problems) >= len(rows), f"{len(problems)} problems")
        check(f"{key}: the broken Excel rows {rows}", sorted({p["row"] for p in problems}) == rows,
              str(sorted({p['row'] for p in problems})))
        got = {c.strip() for p in problems for c in p["column"].split(",")}
        check(f"{key}: Excel column names, not database names", columns <= got, str(got))
        check(f"{key}: nothing loaded, the good file still current",
              body.get("rows_loaded") in (None, 0) and current_file(key) == good_ids.get(key))

    # ------------------------------------------------ the reject list endpoint
    print("\nGET /ingest/files/{file_id}/rejects")
    if "shelf_life" in rejected:
        status, _, raw = request("GET", f"{base}/ingest/files/{rejected['shelf_life']}/rejects")
        body = as_json(raw)
        check("rejected file -> 200 with its problems",
              status == 200 and isinstance(body, dict) and body.get("status") == "rejected" and len(body.get("problems", [])) == 4,
              f"status {status}: {str(body)[:200]}")
        if isinstance(body, dict) and body.get("problems"):
            rows_sorted = [p["row"] for p in body["problems"]]
            check("problems in row order", rows_sorted == sorted(rows_sorted))
            check("summary present", bool(body.get("summary")))
            print("       e.g.", body["problems"][0])
    if "shelf_life" in good_ids:
        status, _, raw = request("GET", f"{base}/ingest/files/{good_ids['shelf_life']}/rejects")
        body = as_json(raw)
        check("published file -> 200 with no problems",
              status == 200 and isinstance(body, dict) and body.get("problems") == [], f"status {status}: {str(body)[:160]}")
    status, _, _ = request("GET", f"{base}/ingest/files/999999/rejects")
    check("unknown file_id -> 404", status == 404, f"status {status}")
    status, _, _ = request("GET", f"{base}/ingest/files/abc/rejects")
    check("file_id not a number -> 422", status == 422, f"status {status}")

    # ------------------------------------------------ the end state
    print("\nend state")
    per_type = dict(db("SELECT file_type, count(*) FROM ingest_files WHERE is_current GROUP BY 1"))
    check("exactly one current upload per type", per_type == {k: 1 for k in FILES}, str(per_type))
    counts = {t: db(f"SELECT count(*) FROM raw_{t} r JOIN ingest_files f USING (file_id) WHERE f.is_current")[0][0]
              for t in FILES}
    check("current rows = the correct files", counts == expected_rows, str(counts))
    if a.reset:                                  # only exact when the test started from an empty folder
        saved = len([f for f in UPLOAD_DIR.glob("*") if f.is_file()])
        uploads = db("SELECT count(*) FROM ingest_files")[0][0]
        check("no orphan files: one saved file per recorded upload", saved == uploads,
              f"{saved} files in uploads/, {uploads} uploads recorded")

    failed = [r for r in results if not r[0]]
    print(f"\n{len(results) - len(failed)} passed, {len(failed)} failed")
    for _, name, detail in failed:
        print(f"  FAIL {name}: {detail}")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
