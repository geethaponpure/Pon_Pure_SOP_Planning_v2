# this file is use for
# Build the downloadable Excel template of every file type from its spec, so a template matches
# what the upload expects. sheet 1 is the template the user fills, sheet 2 the guidelines.
# created in Backend/template/ at api start when missing. after changing a spec, rebuild with:
#   python -m app.ingest.templates --force

import math
from pathlib import Path
from typing import Any, cast

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.ingest.registry import FILE_SPECS
from app.ingest.spec import Col, FileSpec
from app.ingest.validate import _CMP_RULE


TEMPLATE_DIR = Path(__file__).resolve().parents[2] / "template"      # Backend/template
FORMAT_ROWS = 1000                                # rows pre-formatted and checked on the template sheet

# the look agreed on 2026-09-26 (hand-tuned in excel, then built in here)
FONT = "Calibri"
REQUIRED_FILL = PatternFill("solid", fgColor="FF1F4E78")
OPTIONAL_FILL = PatternFill("solid", fgColor="FFD9E1F2")
TITLE_FILL = PatternFill("solid", fgColor=Color(theme=9, tint=0.5999938962981048))     # light green band
SECTION_FILL = PatternFill("solid", fgColor=Color(theme=3, tint=0.5999938962981048))   # light blue-grey band
BULLET_FILL = PatternFill("solid", fgColor=Color(theme=0, tint=-0.1499984740745262))   # light grey
THIN = Side(style="thin")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TITLE = Font(name=FONT, bold=True, size=14)
BOLD = Font(name=FONT, bold=True, size=11)
PLAIN = Font(name=FONT, size=11)
WRAP = Alignment(wrap_text=True, vertical="top")
HEADER_FONT_SIZE = 10
LINE_HEIGHT = 14.5                                # points per wrapped line at 11pt
GUIDE_WIDTHS = {"A": 40, "B": 11, "C": 22, "D": 24, "E": 60, "F": 28}

TYPE_WORDS = {"str": "Text", "int": "Whole number", "float": "Number", "date": "Date (e.g. 25-Sep-2026)",
              "datetime": "Date and time", "bool": "Yes / No"}
NUMBER_FORMATS = {"str": "@", "int": "0", "date": "dd-mmm-yyyy", "datetime": "dd-mmm-yyyy hh:mm"}
RULE_WORDS = {">": "greater than", ">=": "at least", "<": "less than", "<=": "at most", "=": "equal to", "!=": "not"}
EXCEL_OPERATORS = {">": "greaterThan", ">=": "greaterThanOrEqual", "<": "lessThan", "<=": "lessThanOrEqual",
                   "=": "equal", "!=": "notEqual"}


def template_path(key: str) -> Path:
    return TEMPLATE_DIR / f"{key}_template.xlsx"



def _rule_parts(col: Col) -> tuple[str, str] | None:
    """('>=', '0') for a numeric rule, None otherwise."""
    if not col.rule or col.rule.lstrip().startswith("~"):
        return None
    m = _CMP_RULE.match(col.rule)                 # the same syntax validate accepts
    return (m.group(1), m.group(2)) if m else None


def _allowed(col: Col) -> str:
    parts = []
    rule = _rule_parts(col)
    if rule:
        parts.append(f"{RULE_WORDS[rule[0]]} {rule[1]}")
    elif col.rule:
        parts.append(f"must match {col.rule.lstrip()[1:].strip()}")
    if col.null_values:
        parts.append(" or ".join(f"'{v}'" for v in col.null_values) + " for empty")
    return "; ".join(parts)


def _how_it_works(spec: FileSpec) -> list[str]:
    sheet = f"'{spec.sheet}'" if spec.sheet else "the first sheet"
    key_words = ", ".join(next(c.source for c in spec.columns if c.target == k) for k in spec.row_key)

    if spec.scope_column:
        scope = next(c.source for c in spec.columns if c.target == spec.scope_column)
        mode = (f"One {scope} per file: every row must have the same {scope}. Your upload replaces only the "
                f"previous file of that {scope}; files of other {scope}s are not touched.")
    elif spec.mode == "replace":
        mode = "Every upload is the complete list and replaces the previous upload in full."
    else:
        mode = "Every upload covers one period and replaces only the earlier upload of the same period."

    optional = ("Light blue headers are optional: their cells may be left blank, but the column itself must stay "
                "in the file - a file with a column missing is refused."
                if spec.strict_headers else "Light blue headers are optional.")
    lines = [
        f"Fill {sheet} of this workbook. Keep the header row exactly as it is: do not rename or delete header "
        "cells. Column order does not matter, and extra columns are ignored.",
        f"Dark blue headers are required columns: every row needs a value there. {optional}",
        "Hover over a header to see what to enter. The table below lists every column.",
        "Save the file in Excel as Excel Workbook (.xlsx) and upload that. Not .xls or .csv, and not a file "
        "last saved by another program: formulas would have no results.",
        mode,
        "If any row has a problem (a required cell empty, text in a number column, a value outside the allowed "
        "range) the whole file is rejected, nothing is loaded, and you get the list of Excel rows to fix. "
        "The previous upload stays in use until a correct file is loaded.",
        f"Each row is identified by: {key_words}. Two rows with the same values there but different details "
        "reject the file; if one of them simply has more filled in, that one is kept.",
    ]
    if spec.dedupe_exact:
        lines.append("Rows that are identical in every column are counted once.")
    lines.append("Uploading the exact same file twice is refused.")
    return lines



def build_template(spec: FileSpec) -> Workbook:
    wb = Workbook()

    # ---------------- sheet 1: the template
    ws = wb.active
    assert ws is not None
    ws.title = "Template"                         # every spec reads the first sheet, whatever its name
    assert spec.sheet is None, f"{spec.key}: a template named 'Template' would not match sheet={spec.sheet!r}"
    ws.freeze_panes = "A2"

    header_lines = 1
    for i, col in enumerate(spec.columns, start=1):
        letter = get_column_letter(i)
        cell = ws.cell(1, i, col.source)
        cell.fill = REQUIRED_FILL if col.required else OPTIONAL_FILL
        cell.font = Font(name=FONT, size=HEADER_FONT_SIZE, bold=True, color="FFFFFFFF" if col.required else "FF000000")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BOX
        if col.required:
            kind = "Required. "
        elif spec.strict_headers:
            kind = "Optional: cells may be blank, but keep this column. "
        else:
            kind = "Optional. "
        note = kind + (col.help or "")
        if _allowed(col):
            note += f" Allowed: {_allowed(col)}."
        cell.comment = Comment(note.strip(), "SOP tool", width=300, height=110)
        width = min(max(len(col.source) + 4, 14), 42)                 # the header fits on one line where it can
        ws.column_dimensions[letter].width = width
        header_lines = max(header_lines, math.ceil(len(col.source) / (width - 3)))

        fmt = NUMBER_FORMATS.get(col.dtype)
        if fmt:
            for r in range(2, FORMAT_ROWS + 2):
                ws.cell(r, i).number_format = fmt

        # excel's own check, so a typo is caught while typing. skipped where text like '-' is allowed
        if col.dtype in ("int", "float") and not col.null_values:
            rule = _rule_parts(col)
            kind = "whole" if col.dtype == "int" else "decimal"
            if rule:
                dv = DataValidation(type=kind, operator=cast(Any, EXCEL_OPERATORS[rule[0]]), formula1=rule[1],
                                    allow_blank=True)
                dv.error = f"{col.source}: a {TYPE_WORDS[col.dtype].lower()} {RULE_WORDS[rule[0]]} {rule[1]}."
            else:
                dv = DataValidation(type=kind, operator="between", formula1="-1000000000", formula2="1000000000",
                                    allow_blank=True)
                dv.error = f"{col.source}: a {TYPE_WORDS[col.dtype].lower()} only, no units or text."
            dv.errorTitle, dv.showErrorMessage = "Not allowed here", True
            ws.add_data_validation(dv)
            dv.add(f"{letter}2:{letter}{FORMAT_ROWS + 1}")
    ws.row_dimensions[1].height = 8 + 14 * header_lines                # 22 for one line

    # ---------------- sheet 2: the guidelines
    gd = wb.create_sheet("Guidelines")
    for letter, width in GUIDE_WIDTHS.items():
        gd.column_dimensions[letter].width = width
    band_width = sum(GUIDE_WIDTHS.values())

    def band(text: str, font: Font, fill: PatternFill, align: Alignment, height: float | None = None) -> None:
        """One line across A:F, merged, filled and boxed."""
        gd.append([text])
        r = gd.max_row
        gd.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        for c in range(1, 7):                                       # a merged range needs every cell boxed
            gd.cell(r, c).border = BOX
            gd.cell(r, c).fill = fill
        gd.cell(r, 1).font = font
        gd.cell(r, 1).alignment = align
        if height:
            gd.row_dimensions[r].height = height

    def bullet(text: str) -> None:
        lines = math.ceil((len(text) + 2) / (band_width * 1.1))      # wide band, proportional font
        band(f"• {text}", PLAIN, BULLET_FILL, WRAP, LINE_HEIGHT * lines if lines > 1 else None)

    band(f"{spec.label} - how to fill the template", TITLE, TITLE_FILL, Alignment(horizontal="center"), 18.5)
    gd.append([])
    band("How the upload works", BOLD, SECTION_FILL, Alignment(horizontal="center"))
    for line in _how_it_works(spec):
        bullet(line)

    if spec.notes:
        gd.append([])
        band("About this file", BOLD, SECTION_FILL, Alignment(horizontal="center"))
        for line in spec.notes:
            bullet(line)

    gd.append([])
    gd.append(["Column", "Required", "Type", "Allowed values", "What to enter", "Example"])
    head = gd.max_row
    for cell in gd[head]:
        cell.font = Font(name=FONT, size=11, bold=True, color="FFFFFFFF")
        cell.fill = REQUIRED_FILL
        cell.border = BOX
    for col in spec.columns:
        example = "" if col.example is None else col.example
        values = [col.source, "Yes" if col.required else "", TYPE_WORDS[col.dtype], _allowed(col), col.help, example]
        gd.append(values)
        r = gd.max_row
        for cell in gd[r]:
            cell.alignment = WRAP
            cell.border = BOX
        gd.cell(r, 6).number_format = "@"
        gd.cell(r, 6).alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        if col.required:
            gd.cell(r, 1).font = BOLD
        # tallest cell decides the row: roughly one line per column width of text
        lines = max(math.ceil(len(str(v)) / (GUIDE_WIDTHS[get_column_letter(c)] * 1.1)) if v != "" else 1
                    for c, v in enumerate(values, start=1))
        if lines > 1:
            gd.row_dimensions[r].height = LINE_HEIGHT * lines

    wb.active = 0
    return wb



def write_templates(folder: Path = TEMPLATE_DIR, overwrite: bool = False) -> list[Path]:
    """Create the template of every registered file type that does not have one yet, and return the
    files created. An existing template is left alone; after a spec change, delete it (or run with
    --force) so it is built again."""
    folder.mkdir(parents=True, exist_ok=True)
    created = []
    for key, spec in FILE_SPECS.items():
        path = folder / f"{key}_template.xlsx"
        if path.exists() and not overwrite:
            continue
        build_template(spec).save(path)
        created.append(path)
    return created


if __name__ == "__main__":
    import sys
    made = write_templates(overwrite="--force" in sys.argv)
    print("\n".join(str(p) for p in made) if made else f"all templates already exist in {TEMPLATE_DIR} (use --force to rebuild)")
